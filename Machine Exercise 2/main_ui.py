# main_ui.py

from flask import Flask, request, jsonify, send_file
import os
import openpyxl
import sys
import os
# Allow main_ui.py to see inside the voice_pipeline folder
sys.path.append(os.path.join(os.path.dirname(__file__), 'voice_pipeline'))
from intent_parser import parse_text
import datetime
import requests
import glob
import vlc

app = Flask(__name__)

# --- Music Player Setup ---
PLAYLIST_DIR = r"C:\Users\King\Documents\1. UP Files\AI 231 - MLOps\Machine Exercises\me2-repo\Machine Exercise 2\playlist"
if not os.path.exists(PLAYLIST_DIR):
    PLAYLIST_DIR = os.path.join(os.path.dirname(__file__), 'playlist')

playlist = sorted(glob.glob(os.path.join(PLAYLIST_DIR, "*.mp3")))
current_track_idx = 0
player = vlc.MediaPlayer()
current_volume = 50
player.audio_set_volume(current_volume)

def play_current_track():
    if playlist:
        player.set_media(vlc.Media(playlist[current_track_idx]))
        player.play()
        return os.path.basename(playlist[current_track_idx])
    return "No songs found"

# Ensure Excel files exist on startup
for file, header in [('alarms.xlsx', 'Alarm'), ('reminders.xlsx', 'Reminder')]:
    if not os.path.exists(file):
        wb = openpyxl.Workbook()
        wb.active.append([header])
        wb.save(file)

def read_excel(filename):
    if not os.path.exists(filename): return []
    return [str(row[0]) for row in openpyxl.load_workbook(filename).active.iter_rows(min_row=2, values_only=True) if row[0]]

def append_excel(filename, value):
    wb = openpyxl.load_workbook(filename)
    wb.active.append([value])
    wb.save(filename)

@app.route('/get_data', methods=['GET'])
def get_data():
    # Sends saved data to the widgets when the app loads
    return jsonify({"alarms": read_excel('alarms.xlsx'), "reminders": read_excel('reminders.xlsx')})

@app.route('/')
def home():
    return send_file('index.html')

@app.route('/parse_text', methods=['POST'])
def handle_text_parsing():
    data = request.json
    text = data.get('text', '')
    
    # Send text to intent_parser
    action, value = parse_text(text)
    
    return jsonify({"action": action, "value": value})

@app.route('/ringtone')
def get_ringtone():
    return send_file(r"C:\Users\King\Documents\1. UP Files\AI 231 - MLOps\Machine Exercises\me2-repo\Machine Exercise 2\sounds\call_ringtone.mp3")

@app.route('/command', methods=['POST'])
def handle_command():
    global current_track_idx, current_volume
    data = request.json
    action = data.get('action')
    value = data.get('value')

    # IoT Controls
    if action == 'toggle':
        return jsonify({"status": "success", "message": "System: Switched lights."})
    elif action == 'dim':
        if value is not None and 0 <= value <= 100:
            return jsonify({"status": "success", "message": f"System: Dimmed lights to {value}%."})
        return jsonify({"status": "error"})
    elif action == 'color':
        display_color = "default" if value == "none" else value
        return jsonify({
            "status": "success", 
            "message": f"System: Lights adjust to color {display_color}.",
            "value": value,
            "display_color": display_color
        })
    elif action == 'temp':
        return jsonify({"status": "success", "message": f"System: Temperature set to {value} degrees."})
        
    # Media Controls
    elif action == 'music_play':
        if not player.is_playing():
            if player.get_media() is None:
                song_name = play_current_track()
            else:
                player.play()
                song_name = os.path.basename(playlist[current_track_idx])
            return jsonify({"status": "success", "message": f"System: Playing {song_name}", "song": song_name})
        return jsonify({"status": "success", "message": "System: Already playing."})

    elif action == 'music_pause':
        player.pause()
        song_name = os.path.basename(playlist[current_track_idx]) if playlist else "None"
        return jsonify({"status": "success", "message": f"System: Paused {song_name}", "song": song_name})

    elif action == 'music_stop':
        player.stop()
        return jsonify({"status": "success", "message": "System: Music stopped.", "song": "Stopped"})

    elif action == 'music_skip':
        if playlist:
            current_track_idx = (current_track_idx + 1) % len(playlist)
            song_name = play_current_track()
            return jsonify({"status": "success", "message": f"System: Skipped to {song_name}", "song": song_name})
        return jsonify({"status": "error", "message": "System: Playlist empty."})

    elif action == 'music_prev':
        if playlist:
            current_track_idx = (current_track_idx - 1) % len(playlist)
            song_name = play_current_track()
            return jsonify({"status": "success", "message": f"System: Returned to {song_name}", "song": song_name})
        return jsonify({"status": "error", "message": "System: Playlist empty."})

    elif action == 'music_vol_up':
        current_volume = min(100, current_volume + 20)
        player.audio_set_volume(current_volume)
        return jsonify({"status": "success", "message": f"System: Volume at {current_volume}%."})

    elif action == 'music_vol_down':
        current_volume = max(0, current_volume - 20)
        player.audio_set_volume(current_volume)
        return jsonify({"status": "success", "message": f"System: Volume at {current_volume}%."})

    elif action == 'music_vol_set':
        if value is not None and 0 <= value <= 100:
            current_volume = value
            player.audio_set_volume(current_volume)
            return jsonify({"status": "success", "message": f"System: Volume set to {current_volume}%."})
        return jsonify({"status": "error", "message": "System: Invalid volume value."})

    # Search & Questions
    elif action == 'weather':
        try:
            res = requests.get("https://wttr.in/Quezon+City?format=%t,+%C", timeout=5)
            return jsonify({"status": "success", "message": f"Weather: {res.text.strip()}"})
        except requests.RequestException:
            return jsonify({"status": "error", "message": "System: Weather API unreachable."})
            
    elif action == 'time':
        current_time = datetime.datetime.now().strftime("%I:%M %p")
        return jsonify({"status": "success", "message": f"Current Time: {current_time}"})

    # Alarms & Timers
    elif action == 'timer':
        m = value.get("minutes", 0)
        s = value.get("seconds", 0)
        return jsonify({
            "status": "success", 
            "message": f"System: Timer set for {m} minutes and {s} seconds.",
            "minutes": m,
            "seconds": s
        })
        
    elif action == 'alarm':
        append_excel('alarms.xlsx', value)
        return jsonify({"status": "success", "message": f"System: Alarm set for {value}."})

    # Productivity & Calls
    elif action == 'reminder':
        append_excel('reminders.xlsx', value)
        return jsonify({"status": "success", "message": f"System: Reminder added to '{value}'."})
    elif action == 'get_reminders':
        reminders = read_excel('reminders.xlsx')
        text = "\n".join([f"• {r}" for r in reminders]) if reminders else "No reminders."
        return jsonify({"status": "success", "message": "System: Displaying reminders.", "data": text})
    elif action == 'call_mom':
        return jsonify({"status": "success", "message": "System: Calling Mom..."})
    elif action == 'call_dad':
        return jsonify({"status": "success", "message": "System: Calling Dad..."})

    return jsonify({"status": "error"})

# 3-Second Recording
@app.route('/manual_record', methods=['POST'])
def trigger_manual_record():
    # Write a temporary flag file that the wake_word_engine.py will detect
    with open('manual_trigger.txt', 'w') as f:
        f.write('trigger')
    return jsonify({"status": "success", "message": "System: Manual record triggered."})

# --- VOICE RESULT RELAY (engine -> UI) ---
voice_result = {"id": 0, "text": "", "action": None, "value": None}

@app.route('/voice_result', methods=['POST'])
def post_voice_result():
    global voice_result
    text = request.json.get('text', '')
    action, value = parse_text(text)
    voice_result = {"id": voice_result["id"] + 1, "text": text,
                    "action": action, "value": value}
    return jsonify({"status": "success"})

@app.route('/get_voice_result', methods=['GET'])
def get_voice_result():
    return jsonify(voice_result)

# --- SYSTEM STATE TRACKER ---
system_listening_status = "listening"

@app.route('/set_status', methods=['POST'])
def set_status():
    global system_listening_status
    system_listening_status = request.json.get('status', 'listening')
    return jsonify({"status": "success"})

@app.route('/get_status', methods=['GET'])
def get_status():
    return jsonify({"status": system_listening_status})

if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5000)