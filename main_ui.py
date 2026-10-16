# main_ui.py

from flask import Flask, request, jsonify, send_file, Response
import os
import json
import queue
import threading
import time
import openpyxl
import sys
# Allow main_ui.py to see inside the voice_pipeline folder
sys.path.append(os.path.join(os.path.dirname(__file__), 'voice_pipeline'))
from intent_parser import parse_text, command_guide
import datetime
import requests
import glob
import vlc
from collections import deque
import pi_metrics

app = Flask(__name__)

# --- Music Player Setup ---
PLAYLIST_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'playlist')

playlist = sorted(glob.glob(os.path.join(PLAYLIST_DIR, "*.mp3")))
current_track_idx = 0
player = vlc.MediaPlayer()
current_volume = 50
player.audio_set_volume(current_volume)

def play_current_track():
    if playlist:
        player.set_media(vlc.Media(playlist[current_track_idx]))
        player.play()
        return os.path.basename(playlist[current_track_idx])
    return "No songs found"

def current_song():
    return os.path.basename(playlist[current_track_idx]) if playlist else None

def music_state():
    """What the music widget shows: the player's real state, not the browser's guess."""
    state = player.get_state()
    if state == vlc.State.Paused:
        status = "paused"
    elif player.is_playing() or state in (vlc.State.Opening, vlc.State.Buffering):
        status = "playing"
    else:
        status = "stopped"
    return {"status": status, "song": current_song() if status != "stopped" else None, "volume": current_volume}

# Music ducked to DUCK_VOLUME while a voice command is being recorded (restored afterwards)
DUCK_VOLUME = 20
DUCK_FADE_SEC = 0.6
music_ducked = False
_fade_id = 0

def fade_volume(target):
    """Fades the player's volume to target over DUCK_FADE_SEC (a newer fade cancels this one)."""
    global _fade_id
    _fade_id += 1
    my_id = _fade_id
    def run():
        start = player.audio_get_volume()
        if start < 0:
            start = target
        steps = 30
        for i in range(1, steps + 1):
            if _fade_id != my_id:
                return
            player.audio_set_volume(round(start + (target - start) * i / steps))
            time.sleep(DUCK_FADE_SEC / steps)
    threading.Thread(target=run, daemon=True).start()

# --- Push updates to the page (Server-Sent Events) instead of having it poll ---
_listeners = []
_listeners_lock = threading.Lock()

def publish(event, data):
    msg = f"event: {event}\ndata: {json.dumps(data)}\n\n"
    with _listeners_lock:
        for q in list(_listeners):
            q.put(msg)

def _watch_music():
    """Tells the page when the player changes by itself (song ended, etc.)."""
    last = None
    while True:
        try:
            state = music_state()
            if state != last:
                last = state
                publish('music', state)
        except Exception:
            pass
        time.sleep(1)

threading.Thread(target=_watch_music, daemon=True).start()

# --- Pi health readings for the page (only while a page is connected) ---
METRICS_INTERVAL_SEC = 2
metrics_history = deque(maxlen=150)   # last 5 minutes

def _watch_metrics():
    collector = pi_metrics.Collector()
    while True:
        time.sleep(METRICS_INTERVAL_SEC)
        with _listeners_lock:
            watched = bool(_listeners)
        if not watched:
            continue
        try:
            sample = collector.sample()
        except Exception as e:
            print(f"[Metrics] reading failed: {e!r}")
            continue
        metrics_history.append(sample)
        publish('metrics', sample)

if pi_metrics.available():
    threading.Thread(target=_watch_metrics, daemon=True).start()

# --- Weather: fetched in the background, so the command answers at once ---
WEATHER_URL = "https://wttr.in/Quezon+City?format=%t,+%C"
WEATHER_REFRESH_SEC = 600      # refresh every 10 minutes
WEATHER_STALE_SEC = 1800       # older than this: answer anyway, refresh right away
weather_cache = {"text": None, "at": None}
_weather_wakeup = threading.Event()

def _weather_loop():
    while True:
        try:
            res = requests.get(WEATHER_URL, timeout=10)
            res.raise_for_status()
            weather_cache.update(text=res.text.strip(), at=datetime.datetime.now())
            wait = WEATHER_REFRESH_SEC
        except requests.RequestException:
            wait = 60          # offline: try again in a minute
        _weather_wakeup.wait(wait)
        _weather_wakeup.clear()

threading.Thread(target=_weather_loop, daemon=True).start()

def weather_message():
    if weather_cache["text"] is None:
        _weather_wakeup.set()
        return "System: Weather unavailable right now."
    age = (datetime.datetime.now() - weather_cache["at"]).total_seconds()
    if age > WEATHER_STALE_SEC:
        _weather_wakeup.set()
    return f"Weather: {weather_cache['text']} (as of {weather_cache['at'].strftime('%I:%M %p').lstrip('0')})"

# Ensure Excel files exist on startup
for file, header in [('alarms.xlsx', 'Alarm'), ('reminders.xlsx', 'Reminder')]:
    if not os.path.exists(file):
        wb = openpyxl.Workbook()
        wb.active.append([header])
        wb.save(file)

def read_excel(filename):
    if not os.path.exists(filename): return []
    return [str(row[0]) for row in openpyxl.load_workbook(filename).active.iter_rows(min_row=2, values_only=True) if row[0]]

def append_excel(filename, value):
    wb = openpyxl.load_workbook(filename)
    wb.active.append([value])
    wb.save(filename)

@app.route('/get_data', methods=['GET'])
def get_data():
    # Sends saved data to the widgets when the app loads
    return jsonify({"alarms": read_excel('alarms.xlsx'), "reminders": read_excel('reminders.xlsx')})

@app.route('/')
def home():
    return send_file('index.html')

@app.route('/parse_text', methods=['POST'])
def handle_text_parsing():
    data = request.json
    text = data.get('text', '')
    
    # Send text to intent_parser
    action, value = parse_text(text)
    
    return jsonify({"action": action, "value": value})

@app.route('/ringtone')
def get_ringtone():
    return send_file(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'sounds', 'call_ringtone.mp3'))

@app.route('/commands', methods=['GET'])
def get_commands():
    # Phrases the parser understands, grouped by action (for suggestions and the cheat sheet)
    return jsonify(command_guide())

@app.route('/command', methods=['POST'])
def handle_command():
    data = request.json
    return jsonify(execute_command(data.get('action'), data.get('value')))

def execute_command(action, value):
    """Runs one command and returns the response dict for the page."""
    global current_track_idx, current_volume

    # IoT Controls
    if action == 'toggle':
        if value == 'on':
            return ({"status": "success", "message": "System: Lights on."})
        if value == 'off':
            return ({"status": "success", "message": "System: Lights off."})
        return ({"status": "success", "message": "System: Switched lights."})
    elif action == 'dim':
        if value is not None and 0 <= value <= 100:
            return ({"status": "success", "message": f"System: Dimmed lights to {value}%."})
        return ({"status": "error"})
    elif action == 'color':
        display_color = "fire" if value in (None, "none", "fire") else value
        return ({
            "status": "success", 
            "message": f"System: Lights adjust to color {display_color}.",
            "value": value,
            "display_color": display_color
        })
    elif action == 'temp':
        return ({"status": "success", "message": f"System: Temperature set to {value} degrees."})
        
    # Media Controls
    elif action == 'music_play':
        if not playlist:
            return ({"status": "success", "message": "System: No songs in the playlist folder.", "music": music_state()})
        if not player.is_playing():
            if player.get_media() is None:
                song_name = play_current_track()
            else:
                player.play()
                song_name = current_song()
            return ({"status": "success", "message": f"System: Playing {song_name}", "song": song_name, "music": music_state()})
        return ({"status": "success", "message": "System: Already playing.", "music": music_state()})

    elif action == 'music_pause':
        # set_pause(1) always pauses; pause() would toggle and resume a paused song
        if player.is_playing():
            player.set_pause(1)
            return ({"status": "success", "message": f"System: Paused {current_song()}", "song": current_song(), "music": music_state()})
        if player.get_state() == vlc.State.Paused:
            return ({"status": "success", "message": f"System: Paused {current_song()}", "song": current_song(), "music": music_state()})
        return ({"status": "success", "message": "System: Nothing is playing.", "music": music_state()})

    elif action == 'music_stop':
        player.stop()
        return ({"status": "success", "message": "System: Music stopped.", "song": "Stopped", "music": music_state()})

    elif action == 'music_skip':
        if playlist:
            current_track_idx = (current_track_idx + 1) % len(playlist)
            song_name = play_current_track()
            return ({"status": "success", "message": f"System: Skipped to {song_name}", "song": song_name, "music": music_state()})
        return ({"status": "error", "message": "System: Playlist empty."})

    elif action == 'music_prev':
        if playlist:
            current_track_idx = (current_track_idx - 1) % len(playlist)
            song_name = play_current_track()
            return ({"status": "success", "message": f"System: Returned to {song_name}", "song": song_name, "music": music_state()})
        return ({"status": "error", "message": "System: Playlist empty."})

    elif action == 'music_vol_up':
        current_volume = min(100, current_volume + 20)
        player.audio_set_volume(current_volume)
        return ({"status": "success", "message": f"System: Volume at {current_volume}%.", "music": music_state()})

    elif action == 'music_vol_down':
        current_volume = max(0, current_volume - 20)
        player.audio_set_volume(current_volume)
        return ({"status": "success", "message": f"System: Volume at {current_volume}%.", "music": music_state()})

    elif action == 'music_vol_set':
        if value is not None and 0 <= value <= 100:
            current_volume = value
            player.audio_set_volume(current_volume)
            return ({"status": "success", "message": f"System: Volume set to {current_volume}%.", "music": music_state()})
        return ({"status": "error", "message": "System: Invalid volume value."})

    # Search & Questions
    elif action == 'weather':
        return {"status": "success", "message": weather_message()}

    elif action == 'time':
        current_time = datetime.datetime.now().strftime("%I:%M %p")
        return ({"status": "success", "message": f"Current Time: {current_time}"})

    # Alarms & Timers
    elif action == 'timer':
        m = value.get("minutes", 0)
        s = value.get("seconds", 0)
        return ({
            "status": "success", 
            "message": f"System: Timer set for {m} minutes and {s} seconds.",
            "minutes": m,
            "seconds": s
        })
        
    elif action == 'alarm':
        append_excel('alarms.xlsx', value)
        return ({"status": "success", "message": f"System: Alarm set for {value}."})

    # Productivity & Calls
    elif action == 'reminder':
        append_excel('reminders.xlsx', value)
        return ({"status": "success", "message": f"System: Reminder added to '{value}'."})
    elif action == 'get_reminders':
        reminders = read_excel('reminders.xlsx')
        text = "\n".join([f"• {r}" for r in reminders]) if reminders else "No reminders."
        return ({"status": "success", "message": "System: Displaying reminders.", "data": text})
    elif action == 'call_mom':
        return ({"status": "success", "message": "System: Calling Mom..."})
    elif action == 'call_dad':
        return ({"status": "success", "message": "System: Calling Dad..."})
    elif action == 'message':
        return ({"status": "success", "message": "System: Sending message..."})

    return ({"status": "error"})

# 3-Second Recording
@app.route('/manual_record', methods=['POST'])
def trigger_manual_record():
    # Write a temporary flag file that the wake_word_engine.py will detect
    with open('manual_trigger.txt', 'w') as f:
        f.write('trigger')
    return jsonify({"status": "success", "message": "System: Manual record triggered."})

# --- VOICE RESULT RELAY (engine -> UI) ---
voice_result = {"id": 0, "text": "", "action": None, "value": None}

@app.route('/voice_result', methods=['POST'])
def post_voice_result():
    global voice_result
    data = request.json
    text = data.get('text', '')
    action, value = parse_text(text)
    voice_result = {"id": voice_result["id"] + 1, "text": text,
                    "action": action, "value": value}
    # Optional details from the engine's beam search (older engines send text only)
    for key in ('heard', 'top', 'margin', 'infer_ms', 'decode_ms', 'audio_ms', 'end_reason'):
        if key in data:
            voice_result[key] = data[key]
    # Run the command here and now; the page only shows the result
    if action != "Unrecognized":
        voice_result["response"] = execute_command(action, value)
    publish('voice', voice_result)
    return jsonify({"status": "success"})

@app.route('/get_voice_result', methods=['GET'])
def get_voice_result():
    return jsonify(voice_result)

# --- SYSTEM STATE TRACKER ---
system_listening_status = "listening"

@app.route('/set_status', methods=['POST'])
def set_status():
    global system_listening_status, music_ducked
    system_listening_status = request.json.get('status', 'listening')
    # Turn the music down while a command is recorded, and back up afterwards
    if system_listening_status == 'recording' and player.is_playing() and not music_ducked:
        music_ducked = True
        fade_volume(min(DUCK_VOLUME, current_volume))
    elif system_listening_status == 'listening' and music_ducked:
        music_ducked = False
        fade_volume(current_volume)
    publish('status', status_payload())
    return jsonify({"status": "success"})

def status_payload():
    return {"status": system_listening_status, "voice_id": voice_result["id"], "music": music_state()}

@app.route('/get_status', methods=['GET'])
def get_status():
    return jsonify(status_payload())

@app.route('/metrics', methods=['GET'])
def get_metrics():
    # Fallback for when the page has no event stream
    return jsonify(list(metrics_history)[-1:])

@app.route('/events')
def events():
    """Server-Sent Events: status, voice results and music state, the moment they change."""
    q = queue.Queue()
    with _listeners_lock:
        _listeners.append(q)

    def stream():
        try:
            yield f"event: status\ndata: {json.dumps(status_payload())}\n\n"
            if metrics_history:
                yield f"event: metrics_history\ndata: {json.dumps(list(metrics_history))}\n\n"
            while True:
                try:
                    yield q.get(timeout=15)
                except queue.Empty:
                    yield ": keep-alive\n\n"
        finally:
            with _listeners_lock:
                _listeners.remove(q)

    return Response(stream(), mimetype='text/event-stream', headers={'Cache-Control': 'no-cache', 'X-Accel-Buffering': 'no'})

if __name__ == '__main__':
    # A trigger file left over from an earlier run would make the engine record at once
    if os.path.exists('manual_trigger.txt'):
        os.remove('manual_trigger.txt')
    app.run(host='127.0.0.1', port=5000, threaded=True)
